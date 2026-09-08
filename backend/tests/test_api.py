import io
import time

import docx
import fitz
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def test_health_endpoint():
    resp = client.get("/health")
    assert resp.status_code == 200
    assert resp.json()["status"] == "ok"


def _wait_for_job(job_id: str, timeout_s: float = 5.0) -> dict:
    """With CELERY_TASK_ALWAYS_EAGER=true, the job is actually already
    finished by the time POST /jobs returns — the .delay() call runs the
    task synchronously in-process. This loop exists mostly as a safety
    net (and to mirror how a real caller against a real worker would
    behave), not because it typically needs more than one iteration
    here."""
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        resp = client.get(f"/jobs/{job_id}")
        body = resp.json()
        if body["status"] in ("done", "failed"):
            return body
        time.sleep(0.05)
    raise TimeoutError(f"job {job_id} did not finish in time")


def test_create_and_complete_job_happy_path():
    files = {
        "original": ("a.txt", b"The deadline is 30 August 2026.", "text/plain"),
        "modified": ("b.txt", b"The deadline is 15 September 2026.", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    assert create_resp.status_code == 202
    job = create_resp.json()
    assert job["status"] == "queued"
    assert job["id"]

    final = _wait_for_job(job["id"])
    assert final["status"] == "done"
    assert 0.0 <= final["similarity"] <= 1.0
    assert final["stats"]["replaced"] >= 1
    assert final["warnings"] == []


def test_get_job_returns_404_for_unknown_id():
    resp = client.get("/jobs/does-not-exist")
    assert resp.status_code == 404


def test_job_fails_with_friendly_error_for_unsupported_extension():
    files = {
        "original": ("a.exe", b"binary junk", "application/octet-stream"),
        "modified": ("b.txt", b"hello", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    assert create_resp.status_code == 202  # accepted at upload time...

    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "failed"  # ...but fails during processing
    assert "supported file type" in final["error"]


def test_job_fails_for_non_utf8_text():
    files = {
        "original": ("a.txt", b"\xff\xfe\x00bad", "text/plain"),
        "modified": ("b.txt", b"hello", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "failed"
    assert "UTF-8" in final["error"]


def test_job_across_pdf_and_docx():
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "The budget is 500000 rupees.")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    word_doc = docx.Document()
    word_doc.add_paragraph("The budget is 650000 rupees.")
    docx_buf = io.BytesIO()
    word_doc.save(docx_buf)

    files = {
        "original": ("budget.pdf", pdf_bytes, "application/pdf"),
        "modified": (
            "budget.docx",
            docx_buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "done"
    assert final["stats"]["replaced"] >= 1


def test_job_flags_scanned_pdf_warning():
    """Phase 6: same scenario as Phase 3, but now the blank page goes
    through real OCR (finding nothing, since it's genuinely blank),
    producing two OCR warnings. A later feature (DOCX/PDF formatting
    comparison) adds a third warning here too: PDF vs .txt is a mismatched
    pair for formatting comparison (only one side supports it), so that
    gets explicitly skipped-with-explanation as well."""
    pdf_doc = fitz.open()
    pdf_doc.new_page()  # blank page, no text layer
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    files = {
        "original": ("scanned.pdf", pdf_bytes, "application/pdf"),
        "modified": ("notes.txt", b"some text", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "done"
    assert len(final["warnings"]) == 3
    assert any("OCR" in w for w in final["warnings"])
    assert any("Formatting" in w for w in final["warnings"])
    assert any("OCR" in w for w in final["warnings"])


def test_job_fails_for_password_protected_pdf():
    pdf_doc = fitz.open()
    pdf_doc.new_page()
    buf = io.BytesIO()
    pdf_doc.save(buf, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="secret")
    pdf_doc.close()

    files = {
        "original": ("locked.pdf", buf.getvalue(), "application/pdf"),
        "modified": ("notes.txt", b"some text", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "failed"
    assert "password protected" in final["error"].lower()


def test_upload_rejects_oversized_file():
    huge = b"x" * (21 * 1024 * 1024)  # over the 20MB limit
    files = {
        "original": ("big.txt", huge, "text/plain"),
        "modified": ("b.txt", b"hello", "text/plain"),
    }
    resp = client.post("/jobs", files=files)
    assert resp.status_code == 413


def _make_xlsx_bytes(rows: list[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_job_produces_table_diff_for_xlsx_vs_xlsx():
    original = _make_xlsx_bytes([["Name", "Salary"], ["Alice", 45000]])
    modified = _make_xlsx_bytes([["Name", "Salary"], ["Alice", 52000]])

    files = {
        "original": ("payroll.xlsx", original, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "modified": ("payroll.xlsx", modified, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert final["table_diff"] is not None
    table = final["table_diff"]["tables"][0]
    assert table["name"] == "Employees"
    assert table["stats"]["modified"] == 1


def test_job_produces_table_diff_for_csv_vs_csv():
    original = b"Name,Salary\nAlice,45000\nBob,50000\n"
    modified = b"Name,Salary\nAlice,45000\nBob,50000\nCarol,60000\n"

    files = {
        "original": ("people.csv", original, "text/csv"),
        "modified": ("people.csv", modified, "text/csv"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    table = final["table_diff"]["tables"][0]
    assert table["stats"]["added"] == 1


def test_job_skips_table_diff_for_mismatched_formats_with_warning():
    """CSV vs plain text: both extract fine, but only one has tables — no
    row-level diff should be attempted, and the user should be told why."""
    files = {
        "original": ("data.csv", b"a,b\n1,2\n", "text/csv"),
        "modified": ("notes.txt", b"a b 1 2", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert final["table_diff"] is None
    assert any("skipped" in w.lower() for w in final["warnings"])


def test_job_extracts_pptx_slide_text():
    from pptx import Presentation
    from pptx.util import Inches

    def make_pptx(text: str) -> bytes:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
        box.text_frame.text = text
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    files = {
        "original": ("deck.pptx", make_pptx("Q1 revenue is 500000"), "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        "modified": ("deck.pptx", make_pptx("Q1 revenue is 650000"), "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert final["stats"]["replaced"] >= 1


def test_job_produces_visual_diff_for_pdf_vs_pdf():
    import fitz

    def make_pdf(text: str) -> bytes:
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=20)
        raw = doc.tobytes()
        doc.close()
        return raw

    files = {
        "original": ("report.pdf", make_pdf("Quarterly report v1"), "application/pdf"),
        "modified": ("report.pdf", make_pdf("Quarterly report v2 — revised"), "application/pdf"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert final["visual_diff"] is not None
    assert len(final["visual_diff"]) == 1
    page_diff = final["visual_diff"][0]
    assert page_diff["page"] == 1
    assert 0.0 <= page_diff["similarity"] <= 1.0
    assert page_diff["diff_image_key"].startswith("diffs/")


def test_job_skips_visual_diff_for_non_pdf_pair():
    files = {
        "original": ("a.txt", b"hello", "text/plain"),
        "modified": ("b.txt", b"hello there", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "done"
    assert final["visual_diff"] is None


def test_files_endpoint_serves_a_real_diff_image():
    import fitz

    def make_pdf(text: str) -> bytes:
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), text)
        raw = doc.tobytes()
        doc.close()
        return raw

    files = {
        "original": ("a.pdf", make_pdf("version one"), "application/pdf"),
        "modified": ("b.pdf", make_pdf("version two"), "application/pdf"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    key = final["visual_diff"][0]["diff_image_key"]

    resp = client.get(f"/files/{key}")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "image/png"
    assert resp.content[:8] == b"\x89PNG\r\n\x1a\n"  # real PNG magic bytes


def test_files_endpoint_rejects_upload_keys():
    """The whole point of the allow-list regex: this endpoint must NOT be
    usable to read back an uploaded original document, only generated
    diff images."""
    resp = client.get("/files/uploads/some-real-uuid")
    assert resp.status_code == 403


def test_files_endpoint_rejects_path_traversal():
    resp = client.get("/files/diffs/../../etc/passwd")
    # FastAPI/Starlette normalize ".." in path params before routing in
    # some configurations — either a 403 (rejected by our regex) or a 404
    # (no such route/file) is an acceptable outcome; a 200 leaking
    # arbitrary filesystem content would not be.
    assert resp.status_code in (403, 404)


def test_files_endpoint_rejects_non_png_extension():
    resp = client.get("/files/diffs/something.txt")
    assert resp.status_code == 403


def test_files_endpoint_404s_for_a_wellformed_but_nonexistent_key():
    resp = client.get("/files/diffs/00000000-0000-0000-0000-000000000000.png")
    assert resp.status_code == 404


def test_job_ocrs_standalone_images():
    from PIL import Image, ImageDraw

    def make_text_image(text: str) -> bytes:
        img = Image.new("RGB", (600, 150), color="white")
        ImageDraw.Draw(img).text((20, 50), text, fill="black")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return buf.getvalue()

    files = {
        "original": ("note1.png", make_text_image("Meeting at 3pm"), "image/png"),
        "modified": ("note2.png", make_text_image("Meeting at 5pm"), "image/png"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert any("OCR" in w for w in final["warnings"])


def test_job_produces_semantic_diff():
    files = {
        "original": (
            "policy.txt",
            b"Employees must submit expense reports within 30 days.\n\nAll travel requires manager approval.",
            "text/plain",
        ),
        "modified": (
            "policy.txt",
            b"Employees must submit expense reports within 30 days.\n\nInternational travel requires VP approval.",
            "text/plain",
        ),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert final["semantic_diff"] is not None
    assert final["semantic_diff"]["provider"] == "tfidf"
    assert final["semantic_diff"]["stats"]["exact_match"] == 1  # first paragraph unchanged
    changed = [m for m in final["semantic_diff"]["matches"] if m["type"] != "exact_match"]
    assert len(changed) == 1
    assert changed[0]["similarity"] is not None
    assert changed[0]["confidence"] in ("high", "low")


def test_job_produces_formatting_diff_for_docx():
    import docx

    def make_docx(bold: bool) -> bytes:
        document = docx.Document()
        run = document.add_paragraph().add_run("Important notice text")
        run.bold = bold
        buf = io.BytesIO()
        document.save(buf)
        return buf.getvalue()

    files = {
        "original": ("memo.docx", make_docx(False), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "modified": ("memo.docx", make_docx(True), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])

    assert final["status"] == "done"
    assert final["formatting_diff"] is not None
    assert final["formatting_diff"]["stats"]["changed"] == 1
    change = final["formatting_diff"]["changes"][0]
    assert change["changed_properties"]["bold"] == {"from": False, "to": True}


def test_job_skips_formatting_diff_for_unsupported_format_pair():
    files = {
        "original": ("a.txt", b"hello world", "text/plain"),
        "modified": ("b.txt", b"hello there world", "text/plain"),
    }
    create_resp = client.post("/jobs", files=files)
    final = _wait_for_job(create_resp.json()["id"])
    assert final["status"] == "done"
    assert final["formatting_diff"] is None


def test_create_text_job_happy_path():
    payload = {
        "original_text": "Clause 1: Delivery within 30 days.",
        "modified_text": "Clause 1: Delivery within 14 days.",
    }
    resp = client.post("/jobs/text", json=payload)
    assert resp.status_code == 202
    body = resp.json()
    assert body["status"] == "queued"
    assert body["id"]


def test_create_text_job_oversized_input_rejection():
    # 100,001 characters on original side
    oversized_original = {
        "original_text": "x" * 100_001,
        "modified_text": "normal text",
    }
    resp = client.post("/jobs/text", json=oversized_original)
    assert resp.status_code == 413
    assert "Original text exceeds the 100,000 character limit" in resp.json()["detail"]

    # 100,001 characters on modified side
    oversized_modified = {
        "original_text": "normal text",
        "modified_text": "y" * 100_001,
    }
    resp = client.post("/jobs/text", json=oversized_modified)
    assert resp.status_code == 413
    assert "Modified text exceeds the 100,000 character limit" in resp.json()["detail"]


def test_create_text_job_empty_string_handling():
    # Both empty strings
    resp = client.post("/jobs/text", json={"original_text": "", "modified_text": ""})
    assert resp.status_code == 202
    final = _wait_for_job(resp.json()["id"])
    assert final["status"] == "done"
    assert final["similarity"] == 1.0
    assert final["segments"] == []
    assert final["formatting_diff"] is None
    assert final["visual_diff"] is None

    # One empty string (original empty, modified has content)
    resp = client.post(
        "/jobs/text",
        json={"original_text": "", "modified_text": "Freshly added content."},
    )
    assert resp.status_code == 202
    final = _wait_for_job(resp.json()["id"])
    assert final["status"] == "done"
    assert final["similarity"] == 0.0
    assert any(s["type"] == "added" for s in final["segments"])


def test_create_text_job_end_to_end_diff_and_none_fields():
    payload = {
        "original_text": "The quick brown fox jumps over the lazy dog.",
        "modified_text": "The fast brown fox leaps over the lazy dog.",
    }
    resp = client.post("/jobs/text", json=payload)
    assert resp.status_code == 202

    final = _wait_for_job(resp.json()["id"])
    assert final["status"] == "done"
    assert 0.0 < final["similarity"] < 1.0
    assert len(final["segments"]) > 0

    # Verify replaced segments and equal segments
    types = [s["type"] for s in final["segments"]]
    assert "equal" in types
    assert "replaced" in types

    # Paste-to-compare plain text must not have formatting, visual, or table diffs
    assert final["formatting_diff"] is None
    assert final["visual_diff"] is None
    assert final["table_diff"] is None


def test_create_text_job_validation_error():
    # Missing modified_text
    resp = client.post("/jobs/text", json={"original_text": "Hello"})
    assert resp.status_code == 422

