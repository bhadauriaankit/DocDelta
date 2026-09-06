from __future__ import annotations

import io

from pptx import Presentation

from app.parsers.errors import DocumentParseError
from app.parsers.models import TableData


def parse_pptx(raw: bytes, filename: str = "document") -> tuple[str, list[TableData], list[str]]:
    """Extract text from every slide (titles, body placeholders, any text
    box), plus any tables embedded in slides as separate TableData —
    tables inside a deck (a summary table, a comparison grid) get the
    same row-level diffing as a spreadsheet, not just flattened text.
    """
    try:
        presentation = Presentation(io.BytesIO(raw))
    except Exception as exc:
        # Same reasoning as the xlsx/docx parsers: python-pptx's exception
        # types for "not a real pptx" aren't stable enough to enumerate
        # exhaustively, so catch broadly and give one clear message.
        raise DocumentParseError(
            f"'{filename}' couldn't be opened as a PowerPoint file. It may be corrupted or password protected."
        ) from exc

    text_parts: list[str] = []
    tables: list[TableData] = []

    for i, slide in enumerate(presentation.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        slide_lines.append(line)
            if shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                tables.append(TableData(name=f"Slide {i} table", rows=rows))
                for row in rows:
                    slide_lines.append(" | ".join(row))
        if slide_lines:
            text_parts.append(f"Slide {i}:\n" + "\n".join(slide_lines))

    warnings: list[str] = []
    if not text_parts and not tables:
        warnings.append(f"'{filename}' has no extractable text on any slide.")

    return "\n\n".join(text_parts), tables, warnings
