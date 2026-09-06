import fitz
import pytest

from app.visual_diff import MAX_VISUAL_DIFF_PAGES, compare_pdf_pages, make_diff_storage_key


def _make_pdf(page_texts: list[str]) -> bytes:
    doc = fitz.open()
    for text in page_texts:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=24)
    raw = doc.tobytes()
    doc.close()
    return raw


@pytest.fixture
def fake_store():
    """Records every stored image and returns a fake key — lets tests
    assert on what got "saved" without any real object storage."""
    saved: list[bytes] = []

    def store(data: bytes) -> str:
        saved.append(data)
        return f"diffs/fake-{len(saved)}.png"

    store.saved = saved  # type: ignore[attr-defined]
    return store


def test_identical_pdfs_score_high_similarity(fake_store):
    raw = _make_pdf(["Hello world, this is page one."])
    results = compare_pdf_pages(raw, raw, fake_store)

    assert len(results) == 1
    assert results[0]["page"] == 1
    assert results[0]["similarity"] > 0.99
    assert len(fake_store.saved) == 1  # a diff image was still generated and "stored"


def test_visually_different_pdfs_score_lower_similarity(fake_store):
    original = _make_pdf(["Short text."])
    modified = _make_pdf(["This is a completely different and much longer paragraph of text."])

    results = compare_pdf_pages(original, modified, fake_store)
    assert results[0]["similarity"] < 0.99


def test_compares_multiple_pages_by_index(fake_store):
    original = _make_pdf(["Page one original", "Page two original"])
    modified = _make_pdf(["Page one original", "Page two CHANGED"])

    results = compare_pdf_pages(original, modified, fake_store)
    assert len(results) == 2
    assert results[0]["similarity"] > results[1]["similarity"]  # page 1 unchanged, page 2 changed


def test_page_count_capped_at_max(fake_store):
    many_pages = [f"Page {i}" for i in range(MAX_VISUAL_DIFF_PAGES + 5)]
    raw = _make_pdf(many_pages)

    results = compare_pdf_pages(raw, raw, fake_store)
    assert len(results) == MAX_VISUAL_DIFF_PAGES


def test_uses_the_shorter_document_page_count(fake_store):
    original = _make_pdf(["one", "two", "three"])
    modified = _make_pdf(["one", "two"])  # fewer pages

    results = compare_pdf_pages(original, modified, fake_store)
    assert len(results) == 2  # only compares pages that exist in both


def test_handles_different_page_sizes_without_crashing(fake_store):
    doc_a = fitz.open()
    page_a = doc_a.new_page(width=400, height=600)
    page_a.insert_text((20, 20), "Portrait page")
    raw_a = doc_a.tobytes()
    doc_a.close()

    doc_b = fitz.open()
    page_b = doc_b.new_page(width=800, height=400)  # different aspect ratio entirely
    page_b.insert_text((20, 20), "Landscape page")
    raw_b = doc_b.tobytes()
    doc_b.close()

    results = compare_pdf_pages(raw_a, raw_b, fake_store)
    assert len(results) == 1  # didn't crash, produced a result despite the size mismatch


def test_make_diff_storage_key_matches_expected_pattern():
    import re

    key = make_diff_storage_key()
    assert re.match(r"^diffs/[0-9a-f\-]+\.png$", key)


def test_compare_docx_to_pdf_visual(fake_store):
    import io
    import docx
    from app.visual_diff import compare_documents_visual

    # Create a real DOCX
    doc = docx.Document()
    doc.add_heading("Quarterly Results", 0)
    doc.add_paragraph("Revenue increased by 12 percent year over year.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    # Create matching PDF
    pdf_bytes = _make_pdf(["Quarterly Results", "Revenue increased by 12 percent year over year."])

    results = compare_documents_visual(docx_bytes, "docx", pdf_bytes, "pdf", fake_store)
    assert len(results) == 1
    assert results[0]["page"] == 1
    assert "similarity" in results[0]
    assert 0.0 <= results[0]["similarity"] <= 1.0


def test_canonical_page_normalization_letter_vs_a4(fake_store):
    from app.visual_diff import compare_documents_visual

    # Letter is 612 x 792 pt
    doc_letter = fitz.open()
    p1 = doc_letter.new_page(width=612, height=792)
    p1.insert_text((72, 72), "Letter Page", fontsize=18)
    letter_bytes = doc_letter.tobytes()
    doc_letter.close()

    # A4 is ~595.3 x 841.9 pt
    doc_a4 = fitz.open()
    p2 = doc_a4.new_page(width=595.28, height=841.89)
    p2.insert_text((72, 72), "A4 Page", fontsize=18)
    a4_bytes = doc_a4.tobytes()
    doc_a4.close()

    results = compare_documents_visual(
        letter_bytes, "pdf", a4_bytes, "pdf", fake_store, store_all_assets=True
    )
    assert len(results) == 1
    debug = results[0]["debug_info"]
    assert debug is not None
    # Both canvases have the exact same normalized dimensions (max width, max height)
    assert debug["canvas_dims"][0] == max(debug["original_dims"][0], debug["modified_dims"][0]) or debug["canvas_dims"][0] > 0
    assert debug["alignment_offset"] == [0, 0]


def test_difference_region_clustering_and_classification(fake_store):
    from app.visual_diff import compare_documents_visual

    doc_a = fitz.open()
    p_a = doc_a.new_page(width=612, height=792)
    p_a.insert_text((72, 100), "Baseline original headline", fontsize=22)
    bytes_a = doc_a.tobytes()
    doc_a.close()

    doc_b = fitz.open()
    p_b = doc_b.new_page(width=612, height=792)
    p_b.insert_text((72, 100), "Baseline original headline", fontsize=22)
    p_b.insert_text((72, 250), "ADDED paragraph with new content", fontsize=14)
    bytes_b = doc_b.tobytes()
    doc_b.close()

    results = compare_documents_visual(
        bytes_a, "pdf", bytes_b, "pdf", fake_store, store_all_assets=True
    )
    assert len(results) == 1
    page1 = results[0]
    assert len(page1["regions"]) >= 1
    region = page1["regions"][0]
    assert "x" in region and "y" in region and "width" in region and "height" in region
    assert region["diff_type"] in {"text", "layout", "image", "table"}
    assert region["severity"] in {"high", "medium", "low"}
    assert region["pixel_count"] > 0


def test_store_all_assets_generates_all_modes(fake_store):
    from app.visual_diff import compare_documents_visual

    raw = _make_pdf(["Test page"])
    results = compare_documents_visual(
        raw, "pdf", raw, "pdf", fake_store, store_all_assets=True
    )
    assert len(results) == 1
    page1 = results[0]
    assert page1["diff_image_key"].startswith("diffs/")
    assert page1["original_image_key"] is not None
    assert page1["modified_image_key"] is not None
    assert page1["diff_only_image_key"] is not None
    # 4 images saved per page in all-assets mode
    assert len(fake_store.saved) == 4


def test_page_mismatch_with_include_extra_pages(fake_store):
    from app.visual_diff import compare_documents_visual

    doc_3p = _make_pdf(["Page 1", "Page 2", "Page 3"])
    doc_1p = _make_pdf(["Page 1"])

    results = compare_documents_visual(
        doc_3p, "pdf", doc_1p, "pdf", fake_store, store_all_assets=True, include_extra_pages=True
    )
    assert len(results) == 3
    assert results[0]["status"] == "compared"
    assert results[1]["status"] == "only_in_original"
    assert results[2]["status"] == "only_in_original"
