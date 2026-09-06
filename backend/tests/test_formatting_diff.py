import io

import docx
import fitz
from docx.shared import Pt, RGBColor

from app.formatting_diff import (
    compare_formatting,
    extract_formatting_docx,
    extract_formatting_pdf,
)


def _docx_with_run(text: str, **run_kwargs) -> bytes:
    document = docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    for key, value in run_kwargs.items():
        if key == "color":
            run.font.color.rgb = value
        elif key in ("name", "size"):
            setattr(run.font, key, value)
        else:
            setattr(run, key, value)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_formatting_docx_reads_font_and_size():
    document = docx.Document()
    p = document.add_paragraph()
    run = p.add_run("Hello world")
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    buf = io.BytesIO()
    document.save(buf)

    result = extract_formatting_docx(buf.getvalue())
    assert len(result) == 1
    assert result[0].font == "Calibri"
    assert result[0].size == 11.0


def test_extract_formatting_docx_detects_bold_italic_underline():
    document = docx.Document()
    p = document.add_paragraph()
    run = p.add_run("Important text")
    run.bold = True
    run.italic = True
    run.underline = True
    buf = io.BytesIO()
    document.save(buf)

    result = extract_formatting_docx(buf.getvalue())
    assert result[0].bold is True
    assert result[0].italic is True
    assert result[0].underline is True


def test_extract_formatting_docx_reads_color():
    document = docx.Document()
    p = document.add_paragraph()
    run = p.add_run("Red text")
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    buf = io.BytesIO()
    document.save(buf)

    result = extract_formatting_docx(buf.getvalue())
    assert result[0].color == "FF0000"


def test_extract_formatting_docx_skips_blank_paragraphs():
    document = docx.Document()
    document.add_paragraph("")
    document.add_paragraph("Real content")
    buf = io.BytesIO()
    document.save(buf)

    result = extract_formatting_docx(buf.getvalue())
    assert len(result) == 1
    assert result[0].text == "Real content"


def _make_pdf_with_font(text: str, fontname: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontname=fontname, fontsize=12)
    raw = doc.tobytes()
    doc.close()
    return raw


def test_extract_formatting_pdf_detects_bold():
    raw = _make_pdf_with_font("Bold heading", "Helvetica-Bold")
    result = extract_formatting_pdf(raw)
    assert len(result) == 1
    assert result[0].bold is True
    assert result[0].italic is False


def test_extract_formatting_pdf_detects_italic():
    raw = _make_pdf_with_font("Italic note", "Helvetica-Oblique")
    result = extract_formatting_pdf(raw)
    assert result[0].italic is True
    assert result[0].bold is False


def test_extract_formatting_pdf_plain_text_has_no_styling():
    raw = _make_pdf_with_font("Plain paragraph", "Helvetica")
    result = extract_formatting_pdf(raw)
    assert result[0].bold is False
    assert result[0].italic is False


# --- compare_formatting: the matching + diffing logic ---


def test_identical_formatting_reports_unchanged():
    a = extract_formatting_docx(_docx_with_run("Same text", name="Arial"))
    b = extract_formatting_docx(_docx_with_run("Same text", name="Arial"))
    result = compare_formatting(a, b)
    assert result.stats == {"unchanged": 1, "changed": 0, "added": 0, "removed": 0}


def test_font_change_on_same_text_is_reported():
    a = extract_formatting_docx(_docx_with_run("Same text", name="Calibri"))
    b = extract_formatting_docx(_docx_with_run("Same text", name="Arial"))
    result = compare_formatting(a, b)

    assert result.stats["changed"] == 1
    change = result.changes[0]
    assert change.changed_properties["font"] == {"from": "Calibri", "to": "Arial"}


def test_bold_added_on_same_text_is_reported():
    a = extract_formatting_docx(_docx_with_run("Warning text"))
    b = extract_formatting_docx(_docx_with_run("Warning text", bold=True))
    result = compare_formatting(a, b)

    assert result.stats["changed"] == 1
    assert result.changes[0].changed_properties["bold"] == {"from": False, "to": True}


def test_multiple_property_changes_all_reported():
    a = extract_formatting_docx(_docx_with_run("Text", name="Calibri", size=Pt(11)))
    b = extract_formatting_docx(_docx_with_run("Text", name="Arial", size=Pt(14), bold=True))
    result = compare_formatting(a, b)

    props = result.changes[0].changed_properties
    assert "font" in props
    assert "size" in props
    assert "bold" in props


def test_added_and_removed_paragraphs():
    a = extract_formatting_docx(_docx_with_run("Kept paragraph"))
    document = docx.Document()
    document.add_paragraph("Kept paragraph")
    document.add_paragraph("Brand new paragraph")
    buf = io.BytesIO()
    document.save(buf)
    b = extract_formatting_docx(buf.getvalue())

    result = compare_formatting(a, b)
    assert result.stats["unchanged"] == 1
    assert result.stats["added"] == 1


def test_unrelated_paragraphs_are_not_paired():
    """Below the similarity floor, two unrelated paragraphs should show
    up as one removed + one added, NOT as a misleading 'formatting
    changed' pairing between content that isn't really the same
    paragraph."""
    a = extract_formatting_docx(_docx_with_run("Quarterly revenue increased by twelve percent"))
    b = extract_formatting_docx(_docx_with_run("Please remember to water the office plants"))

    result = compare_formatting(a, b)
    assert result.stats["removed"] == 1
    assert result.stats["added"] == 1
    assert result.stats["changed"] == 0


def test_empty_inputs_produce_only_additions_or_removals():
    a = extract_formatting_docx(_docx_with_run("Something"))
    result = compare_formatting(a, [])
    assert result.stats["removed"] == 1

    result2 = compare_formatting([], a)
    assert result2.stats["added"] == 1
