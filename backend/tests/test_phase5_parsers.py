import io

import openpyxl
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.parsers.csv_parser import parse_csv
from app.parsers.errors import DocumentParseError
from app.parsers.pptx_parser import parse_pptx
from app.parsers.xlsx_parser import parse_xlsx


def _make_xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default blank "Sheet"
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_xlsx_extracts_sheets_as_tables():
    raw = _make_xlsx_bytes({"Employees": [["Name", "Salary"], ["Alice", 45000]]})
    text, tables, warnings = parse_xlsx(raw, "data.xlsx")

    assert len(tables) == 1
    assert tables[0].name == "Employees"
    assert tables[0].rows == [["Name", "Salary"], ["Alice", "45000"]]
    assert "Employees" in text
    assert warnings == []


def test_parse_xlsx_preserves_formulas_as_text():
    raw = _make_xlsx_bytes({"Sheet1": [["A", "B", "Total"], [1, 2, "=A1+B1"]]})
    _, tables, _ = parse_xlsx(raw, "formulas.xlsx")
    assert tables[0].rows[1][2] == "=A1+B1"


def test_parse_xlsx_handles_multiple_sheets():
    raw = _make_xlsx_bytes({"Q1": [["100"]], "Q2": [["200"]]})
    _, tables, _ = parse_xlsx(raw, "quarters.xlsx")
    names = [t.name for t in tables]
    assert names == ["Q1", "Q2"]


def test_parse_xlsx_rejects_corrupted_file():
    with pytest.raises(DocumentParseError, match="couldn't be opened as an Excel file"):
        parse_xlsx(b"not a real xlsx at all", "broken.xlsx")


def test_parse_xlsx_warns_on_no_data():
    raw = _make_xlsx_bytes({"Empty": []})
    _, _, warnings = parse_xlsx(raw, "empty.xlsx")
    assert len(warnings) == 1


def test_parse_csv_extracts_rows():
    raw = b"Name,Salary\nAlice,45000\nBob,50000\n"
    text, tables, warnings = parse_csv(raw, "people.csv")

    assert len(tables) == 1
    assert tables[0].name == "csv"
    assert tables[0].rows == [["Name", "Salary"], ["Alice", "45000"], ["Bob", "50000"]]
    assert warnings == []


def test_parse_csv_keeps_values_as_strings_not_numbers():
    """The whole reason we didn't use pandas — "45000" should stay
    exactly "45000", not become 45000 (int) or 45000.0 (float)."""
    raw = b"value\n007\n"
    _, tables, _ = parse_csv(raw, "codes.csv")
    assert tables[0].rows == [["value"], ["007"]]


def test_parse_csv_warns_on_empty_file():
    _, _, warnings = parse_csv(b"", "empty.csv")
    assert len(warnings) == 1


def _make_pptx_bytes(slides_text: list[list[str]], with_table: bool = False) -> bytes:
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank layout, no placeholders to fight with
    for lines in slides_text:
        slide = prs.slides.add_slide(blank_layout)
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
        tf = textbox.text_frame
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line

    if with_table:
        slide = prs.slides.add_slide(blank_layout)
        table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Salary"
        table.cell(1, 0).text = "Alice"
        table.cell(1, 1).text = "45000"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_parse_pptx_extracts_slide_text():
    raw = _make_pptx_bytes([["Welcome", "Agenda for today"], ["Next steps"]])
    text, tables, warnings = parse_pptx(raw, "deck.pptx")

    assert "Welcome" in text
    assert "Agenda for today" in text
    assert "Next steps" in text
    assert text.index("Welcome") < text.index("Next steps")
    assert tables == []
    assert warnings == []


def test_parse_pptx_extracts_embedded_table():
    raw = _make_pptx_bytes([["Intro"]], with_table=True)
    _, tables, _ = parse_pptx(raw, "deck.pptx")

    assert len(tables) == 1
    assert tables[0].rows == [["Name", "Salary"], ["Alice", "45000"]]


def test_parse_pptx_rejects_corrupted_file():
    with pytest.raises(DocumentParseError, match="couldn't be opened as a PowerPoint file"):
        parse_pptx(b"not a real pptx at all", "broken.pptx")
