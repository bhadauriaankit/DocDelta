import io

import docx
import fitz
import pytest

from app.parsers.docx_parser import parse_docx
from app.parsers.errors import DocumentParseError
from app.parsers.pdf_parser import parse_pdf


def _make_pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    raw = doc.tobytes()
    doc.close()
    return raw


def _make_blank_pdf_bytes() -> bytes:
    doc = fitz.open()
    doc.new_page()
    raw = doc.tobytes()
    doc.close()
    return raw


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_parse_pdf_extracts_text():
    raw = _make_pdf_bytes("The deadline is 30 August 2026.")
    text, warnings = parse_pdf(raw, "deadline.pdf")
    assert "deadline" in text.lower()
    assert warnings == []


def test_parse_pdf_warns_on_near_empty_text():
    """Phase 6: a blank/scanned page now goes through OCR rather than
    just getting a "this looks scanned" warning — OCR runs, finds
    nothing (there's genuinely no text on a blank page), and both the
    "OCR was used" and "OCR found nothing" warnings are surfaced."""
    raw = _make_blank_pdf_bytes()
    text, warnings = parse_pdf(raw, "scanned.pdf")
    assert len(warnings) == 2
    assert any("OCR" in w for w in warnings)
    assert any("little to no readable text" in w for w in warnings)


def test_parse_pdf_rejects_corrupted_file():
    with pytest.raises(DocumentParseError, match="couldn't be opened as a PDF"):
        parse_pdf(b"%PDF-1.7 but then garbage", "broken.pdf")


def test_parse_docx_extracts_paragraphs_in_order():
    raw = _make_docx_bytes(["First paragraph.", "Second paragraph."])
    text = parse_docx(raw, "doc.docx")
    assert text.index("First paragraph") < text.index("Second paragraph")


def test_parse_docx_rejects_corrupted_file():
    with pytest.raises(DocumentParseError, match="couldn't be opened as a Word document"):
        parse_docx(b"not a real docx at all", "broken.docx")


def test_parse_docx_includes_table_cells():
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Salary"
    buf = io.BytesIO()
    document.save(buf)

    text = parse_docx(buf.getvalue(), "table.docx")
    assert "Name" in text
    assert "Salary" in text
